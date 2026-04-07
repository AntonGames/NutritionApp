from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..profile import ProfileConfig
from ..services.summaries import SummaryService
from ..storage import Database


HEADER_FILL = PatternFill("solid", fgColor="2F5D50")
WHITE_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=16, bold=True)
SECTION_FONT = Font(size=12, bold=True)
ALERT_FILL = PatternFill("solid", fgColor="F7E7E1")


def _autosize(sheet: Worksheet) -> None:
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 30)


@dataclass(slots=True)
class WorkbookExporter:
    database: Database
    profile: ProfileConfig
    summary_service: SummaryService

    def export(self, target_path: Path) -> Path:
        workbook = Workbook()
        workbook.remove(workbook.active)

        snapshot = self.database.export_snapshot()
        daily_rollups = self.database.build_daily_rollups()
        weekly_rollups = self.database.get_weekly_rollups(self.profile)
        today_summary = self.summary_service.build_daily_summary(datetime.now().date())

        self._build_dashboard_sheet(workbook, daily_rollups, weekly_rollups, today_summary)
        self._build_settings_sheet(workbook)
        self._build_meals_sheet(workbook, snapshot["meals"], snapshot["meal_components"])
        self._build_weights_sheet(workbook, snapshot["weights"])
        self._build_workouts_sheet(workbook, snapshot["workouts"])
        self._build_daily_sheet(workbook, daily_rollups)
        self._build_weekly_sheet(workbook, weekly_rollups)

        target_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target_path)
        return target_path

    def _build_dashboard_sheet(self, workbook: Workbook, daily_rollups: list[dict], weekly_rollups: list[dict], today_summary) -> None:
        sheet = workbook.create_sheet("Главная")
        sheet["A1"] = "Трекер питания и прогресса"
        sheet["A1"].font = TITLE_FONT

        cards = [
            ("Текущий вес", today_summary.weight_kg if today_summary.weight_kg is not None else "—"),
            ("Калории сегодня", today_summary.food.calories),
            ("Осталось калорий", today_summary.remaining.calories),
            ("Белок сегодня", today_summary.food.protein_g),
        ]
        row = 3
        for index, (title, value) in enumerate(cards):
            column = 1 + index * 2
            title_cell = sheet.cell(row=row, column=column, value=title)
            value_cell = sheet.cell(row=row + 1, column=column, value=value)
            title_cell.fill = HEADER_FILL
            title_cell.font = WHITE_FONT
            value_cell.font = Font(size=14, bold=True)

        sheet["A7"] = "Подсказки"
        sheet["A7"].font = SECTION_FONT
        tips = today_summary.suggestions or ["Пока нет рекомендаций."]
        for offset, text in enumerate(tips, start=8):
            sheet[f"A{offset}"] = f"• {text}"

        sheet["F7"] = "Последние 14 дней"
        sheet["F7"].font = SECTION_FONT
        headers = ["Дата", "Вес", "Еда kcal", "Белок", "Тренировки"]
        for column, header in enumerate(headers, start=6):
            cell = sheet.cell(row=8, column=column, value=header)
            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT
        recent_days = daily_rollups[-14:]
        for row_index, item in enumerate(recent_days, start=9):
            sheet.cell(row=row_index, column=6, value=item["date"])
            sheet.cell(row=row_index, column=7, value=item["weight_kg"])
            sheet.cell(row=row_index, column=8, value=item["food_calories"])
            sheet.cell(row=row_index, column=9, value=item["protein_g"])
            sheet.cell(row=row_index, column=10, value=item["workouts_count"])

        sheet["A15"] = "Последние недели"
        sheet["A15"].font = SECTION_FONT
        weekly_headers = ["Неделя", "Средний вес", "Δ", "Средние kcal", "Средний белок", "Тренировки", "Комментарий"]
        for column, header in enumerate(weekly_headers, start=1):
            cell = sheet.cell(row=16, column=column, value=header)
            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT
        for row_index, item in enumerate(weekly_rollups[-8:], start=17):
            sheet.cell(row=row_index, column=1, value=item["week_start"])
            sheet.cell(row=row_index, column=2, value=item["average_weight_kg"])
            sheet.cell(row=row_index, column=3, value=item["weight_delta_vs_previous_week_kg"])
            sheet.cell(row=row_index, column=4, value=item["average_food_calories"])
            sheet.cell(row=row_index, column=5, value=item["average_protein_g"])
            sheet.cell(row=row_index, column=6, value=item["workouts"])
            sheet.cell(row=row_index, column=7, value=item["coach_note"])

        if recent_days:
            line_chart = LineChart()
            line_chart.title = "Вес за последние дни"
            line_chart.height = 7
            line_chart.width = 13
            data = Reference(sheet, min_col=7, min_row=8, max_row=8 + len(recent_days))
            cats = Reference(sheet, min_col=6, min_row=9, max_row=8 + len(recent_days))
            line_chart.add_data(data, titles_from_data=True)
            line_chart.set_categories(cats)
            sheet.add_chart(line_chart, "F18")

            bar_chart = BarChart()
            bar_chart.title = "Калории по дням"
            bar_chart.height = 7
            bar_chart.width = 13
            kcal_data = Reference(sheet, min_col=8, min_row=8, max_row=8 + len(recent_days))
            bar_chart.add_data(kcal_data, titles_from_data=True)
            bar_chart.set_categories(cats)
            sheet.add_chart(bar_chart, "F33")

        _autosize(sheet)

    def _build_settings_sheet(self, workbook: Workbook) -> None:
        sheet = workbook.create_sheet("Настройки")
        sheet["A1"] = "Профиль"
        sheet["A1"].font = TITLE_FONT
        rows = [
            ("Имя", self.profile.user.name),
            ("Часовой пояс", self.profile.timezone),
            ("Рост (см)", self.profile.user.height_cm),
            ("Возраст", self.profile.user.age),
            ("Пол", self.profile.user.sex),
            ("Стартовый вес (кг)", self.profile.user.start_weight_kg),
            ("Целевой вес (кг)", self.profile.user.goal_weight_kg),
            ("Калории", self.profile.targets.daily_calories),
            ("Белок", self.profile.targets.protein_g),
            ("Жиры", self.profile.targets.fat_g),
            ("Углеводы", self.profile.targets.carbs_g),
            ("Белковый минимум", self.profile.preferences.protein_floor_g),
            ("Тренировок в неделю", self.profile.preferences.weekly_workout_goal),
        ]
        for idx, (label, value) in enumerate(rows, start=3):
            sheet[f"A{idx}"] = label
            sheet[f"B{idx}"] = value
        _autosize(sheet)

    def _build_meals_sheet(self, workbook: Workbook, meals: list[dict], components: list[dict]) -> None:
        sheet = workbook.create_sheet("Приемы пищи")
        headers = [
            "Дата",
            "Время",
            "Блюдо",
            "Компонент",
            "Категория",
            "Оценка, г",
            "Ккал",
            "Белки",
            "Жиры",
            "Углеводы",
            "Точность",
            "Источник",
            "Комментарий",
            "Фото",
        ]
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT
        meal_map = {item["id"]: item for item in meals}
        row = 2
        for component in components:
            meal = meal_map[component["meal_id"]]
            values = [
                meal["local_date"],
                meal["local_time"],
                meal["meal_name"],
                component["description"],
                component["category"],
                component["estimated_grams"],
                component["calories"],
                component["protein_g"],
                component["fat_g"],
                component["carbs_g"],
                meal["confidence"],
                meal["source"],
                meal["note"],
                meal["image_path"],
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row, column=column, value=value)
            row += 1
        sheet.freeze_panes = "A2"
        _autosize(sheet)

    def _build_weights_sheet(self, workbook: Workbook, weights: list[dict]) -> None:
        sheet = workbook.create_sheet("Вес")
        headers = ["Дата", "Время", "Вес (кг)", "Источник", "Комментарий"]
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT
        for row, item in enumerate(weights, start=2):
            sheet.cell(row=row, column=1, value=item["local_date"])
            sheet.cell(row=row, column=2, value=item["local_time"])
            sheet.cell(row=row, column=3, value=item["weight_kg"])
            sheet.cell(row=row, column=4, value=item["source"])
            sheet.cell(row=row, column=5, value=item["note"])
        _autosize(sheet)

    def _build_workouts_sheet(self, workbook: Workbook, workouts: list[dict]) -> None:
        sheet = workbook.create_sheet("Тренировки")
        headers = ["Дата", "Время", "Описание", "Минуты", "Ккал", "Средний пульс", "Источник", "Комментарий"]
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT
        for row, item in enumerate(workouts, start=2):
            sheet.cell(row=row, column=1, value=item["local_date"])
            sheet.cell(row=row, column=2, value=item["local_time"])
            sheet.cell(row=row, column=3, value=item["description"])
            sheet.cell(row=row, column=4, value=item["duration_min"])
            sheet.cell(row=row, column=5, value=item["calories_burned"])
            sheet.cell(row=row, column=6, value=item["avg_hr"])
            sheet.cell(row=row, column=7, value=item["source"])
            sheet.cell(row=row, column=8, value=item["note"])
        _autosize(sheet)

    def _build_daily_sheet(self, workbook: Workbook, daily_rollups: list[dict]) -> None:
        sheet = workbook.create_sheet("Дни")
        headers = [
            "Дата",
            "Вес",
            "Ккал еды",
            "Белок",
            "Жиры",
            "Углеводы",
            "Ккал тренировок",
            "Чистые ккал",
            "Приемов пищи",
            "Тренировок",
        ]
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT
        for row, item in enumerate(daily_rollups, start=2):
            values = [
                item["date"],
                item["weight_kg"],
                item["food_calories"],
                item["protein_g"],
                item["fat_g"],
                item["carbs_g"],
                item["exercise_calories"],
                item["net_calories"],
                item["meals_count"],
                item["workouts_count"],
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row, column=column, value=value)
        _autosize(sheet)

    def _build_weekly_sheet(self, workbook: Workbook, weekly_rollups: list[dict]) -> None:
        sheet = workbook.create_sheet("Недели")
        headers = ["Неделя", "Средний вес", "Δ к прошлой", "Средние ккал", "Средний белок", "Тренировки", "Дней", "Комментарий"]
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT
        for row, item in enumerate(weekly_rollups, start=2):
            values = [
                item["week_start"],
                item["average_weight_kg"],
                item["weight_delta_vs_previous_week_kg"],
                item["average_food_calories"],
                item["average_protein_g"],
                item["workouts"],
                item["days_logged"],
                item["coach_note"],
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                if column == 8 and "ниже" in str(value).lower():
                    cell.fill = ALERT_FILL
        _autosize(sheet)
